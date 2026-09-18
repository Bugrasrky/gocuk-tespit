import math
import os
import openpyxl
from paths import resource_path

DEFAULT_SCENARIOS = {
    "OTOPARK": {"name": "Otopark / Kapı Çarpması", "k_min": 0.04, "k_max": 0.08, "multiplier": 1.0, "stretch_risk": "Düşük", "paint_risk": "Çok Düşük"},
    "DOLU_TAS": {"name": "Dolu / Küçük Taş Çarpması", "k_min": 0.03, "k_max": 0.06, "multiplier": 0.9, "stretch_risk": "Yok", "paint_risk": "Yok"},
    "MOTOR_BISIKLET": {"name": "Motosiklet / Bisiklet Teması", "k_min": 0.12, "k_max": 0.20, "multiplier": 1.6, "stretch_risk": "Yüksek", "paint_risk": "Orta-Yüksek"},
    "BARIYER_KOLON": {"name": "Bariyer / Kolon / Duvar Sürtmesi", "k_min": 0.06, "k_max": 0.12, "multiplier": 1.4, "stretch_risk": "Orta", "paint_risk": "Çok Yüksek"},
    "DIREK_CARPMA": {"name": "Geri Manevra / Direk Çarpması", "k_min": 0.10, "k_max": 0.18, "multiplier": 1.5, "stretch_risk": "Yüksek", "paint_risk": "Orta"},
    "TOP_YUMRUK": {"name": "Top / Yumruk / İnsan Teması", "k_min": 0.05, "k_max": 0.09, "multiplier": 1.0, "stretch_risk": "Düşük", "paint_risk": "Yok"}
}

DEFAULT_PRICING = [
    {"class": "Mini (Dolu / Küçük Tıklama)", "min_d": 0.0, "max_d": 3.0, "base_min": 750, "base_max": 1200, "paint_cost": 1500, "duration": "30 - 45 Dk"},
    {"class": "Küçük (Yumruk / Kapı Teması)", "min_d": 3.0, "max_d": 7.0, "base_min": 1300, "base_max": 2200, "paint_cost": 2000, "duration": "1 - 2 Saat"},
    {"class": "Orta (Geniş Darbe)", "min_d": 7.0, "max_d": 12.0, "base_min": 2200, "base_max": 3800, "paint_cost": 2800, "duration": "2 - 4 Saat"},
    {"class": "Büyük (Bariyer / Direk Teması)", "min_d": 12.0, "max_d": 20.0, "base_min": 3800, "base_max": 6500, "paint_cost": 3500, "duration": "1 Gün"},
    {"class": "Çok Büyük (Panel Deformasyonu)", "min_d": 20.0, "max_d": 50.0, "base_min": 6500, "base_max": 11000, "paint_cost": 5000, "duration": "1 - 2 Gün"}
]


class DamageEstimator:
    def __init__(self, excel_filename="hasar_derinlik_matrisi.xlsx"):
        self.excel_path = resource_path(excel_filename)
        self.scenarios = {}
        self.pricing_bands = []
        self._load_database()

    def _load_database(self):
        if not os.path.exists(self.excel_path):
            self.scenarios = DEFAULT_SCENARIOS.copy()
            self.pricing_bands = DEFAULT_PRICING.copy()
            return

        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            ws1 = wb["Hasar_Senaryolari"]
            for row in ws1.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    continue
                s_id, name, k_min, k_max, energy, stretch_risk, paint_risk, multiplier = row[:8]
                self.scenarios[str(s_id)] = {
                    "name": str(name), "k_min": float(k_min), "k_max": float(k_max),
                    "stretch_risk": str(stretch_risk), "paint_risk": str(paint_risk),
                    "multiplier": float(multiplier)
                }

            ws2 = wb["Fiyat_Matrisi"]
            for row in ws2.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    continue
                size_class, min_d, max_d, base_min, base_max, paint_cost, duration = row[:7]
                self.pricing_bands.append({
                    "class": str(size_class), "min_d": float(min_d), "max_d": float(max_d),
                    "base_min": float(base_min), "base_max": float(base_max),
                    "paint_cost": float(paint_cost), "duration": str(duration)
                })
        except Exception:
            self.scenarios = DEFAULT_SCENARIOS.copy()
            self.pricing_bands = DEFAULT_PRICING.copy()

    def get_scenario_list(self):
        return [(s_id, data["name"]) for s_id, data in self.scenarios.items()]

    def estimate(self, length_cm, width_cm, area_cm2, scenario_id, paint_damaged=False):
        scenario = self.scenarios.get(scenario_id, DEFAULT_SCENARIOS["OTOPARK"])

        effective_diameter_cm = math.sqrt(length_cm * width_cm) if width_cm > 0 else length_cm
        effective_diameter_mm = effective_diameter_cm * 10.0

        depth_min_mm = effective_diameter_mm * scenario["k_min"]
        depth_max_mm = effective_diameter_mm * scenario["k_max"]

        if depth_max_mm <= 2.5:
            severity = "Hafif / Sığ (Standart Masaj)"
        elif depth_max_mm <= 6.0:
            severity = "Orta Derinlik"
        else:
            severity = "Derin / Sac Uzamış Olabilir"

        matched_band = None
        for band in self.pricing_bands:
            if band["min_d"] <= effective_diameter_cm < band["max_d"]:
                matched_band = band
                break
        if not matched_band and self.pricing_bands:
            matched_band = self.pricing_bands[-1]

        mult = scenario["multiplier"]
        pdr_min = int(matched_band["base_min"] * mult)
        pdr_max = int(matched_band["base_max"] * mult)

        paint_cost_val = int(matched_band.get("paint_cost", 2000)) if paint_damaged else 0
        total_cost_min = pdr_min + paint_cost_val
        total_cost_max = pdr_max + paint_cost_val

        return {
            "scenario_name": scenario["name"],
            "effective_diameter_cm": effective_diameter_cm,
            "depth_min_mm": depth_min_mm,
            "depth_max_mm": depth_max_mm,
            "severity": severity,
            "stretch_risk": scenario["stretch_risk"],
            "paint_risk": scenario["paint_risk"],
            "paint_damaged": paint_damaged,
            "paint_cost_added": paint_cost_val,
            "cost_min": total_cost_min,
            "cost_max": total_cost_max,
            "duration": matched_band["duration"],
            "size_class": matched_band["class"]
        }